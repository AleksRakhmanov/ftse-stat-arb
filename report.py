import logging

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

import config

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


# ---------------------------------------------------------------------------
# Matplotlib charts (saved as PNG, then embedded in the workbook)
# ---------------------------------------------------------------------------

def plot_pair_chart(signal_frame: pd.DataFrame, daily: pd.DataFrame, ticker_a: str, ticker_b: str, out_path) -> str:
    """3-panel chart: normalised prices, spread z-score with entry/exit bands, equity curve."""
    fig, axes = plt.subplots(3, 1, figsize=(10, 9), sharex=True)

    axes[0].plot(signal_frame.index, signal_frame["price_a"] / signal_frame["price_a"].iloc[0], label=ticker_a)
    axes[0].plot(signal_frame.index, signal_frame["price_b"] / signal_frame["price_b"].iloc[0], label=ticker_b)
    axes[0].set_title(f"{ticker_a} vs {ticker_b} — normalised price")
    axes[0].legend(loc="upper left")
    axes[0].grid(alpha=0.3)

    axes[1].plot(signal_frame.index, signal_frame["zscore"], color="black", linewidth=0.9, label="z-score")
    axes[1].axhline(config.ENTRY_Z, color="red", linestyle="--", linewidth=0.8, label="+entry")
    axes[1].axhline(-config.ENTRY_Z, color="green", linestyle="--", linewidth=0.8, label="-entry")
    axes[1].axhline(config.EXIT_Z, color="grey", linestyle=":", linewidth=0.8, label="exit")
    axes[1].set_title("Spread z-score")
    axes[1].legend(loc="upper left", fontsize=8)
    axes[1].grid(alpha=0.3)

    axes[2].plot(daily.index, daily["portfolio_value"], color="navy")
    axes[2].set_title("Pair portfolio value")
    axes[2].grid(alpha=0.3)

    fig.tight_layout()
    fig.savefig(out_path, dpi=120)
    plt.close(fig)
    return str(out_path)


def plot_portfolio_curve(aggregate_daily: pd.Series, out_path) -> str:
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.plot(aggregate_daily.index, aggregate_daily.values, color="darkgreen")
    ax.set_title("Aggregate strategy portfolio value (all traded pairs)")
    ax.set_ylabel("Portfolio value (GBP)")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(out_path, dpi=120)
    plt.close(fig)
    return str(out_path)


def plot_drawdown(aggregate_daily: pd.Series, out_path) -> str:
    running_max = aggregate_daily.cummax()
    drawdown = aggregate_daily / running_max - 1.0

    fig, ax = plt.subplots(figsize=(10, 3.5))
    ax.fill_between(drawdown.index, drawdown.values * 100, 0, color="firebrick", alpha=0.5)
    ax.set_title("Aggregate strategy drawdown")
    ax.set_ylabel("Drawdown (%)")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(out_path, dpi=120)
    plt.close(fig)
    return str(out_path)


# ---------------------------------------------------------------------------
# Excel dashboard
# ---------------------------------------------------------------------------

def _write_dataframe(ws, df: pd.DataFrame, start_row: int = 1):
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

    for c_idx, col in enumerate(df.columns, start=1):
        max_len = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).values])
        ws.column_dimensions[get_column_letter(c_idx)].width = min(max_len + 2, 40)


def build_excel_dashboard(
    summary_metrics: dict,
    pair_stats: pd.DataFrame,
    trade_log: pd.DataFrame,
    chart_paths: dict,
    out_path,
):
    wb = Workbook()

    # --- Summary sheet -----------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary["A1"] = "FTSE 100 Statistical Arbitrage — Pairs Trading Dashboard"
    ws_summary["A1"].font = Font(bold=True, size=14)

    ws_summary["A3"] = "Metric"
    ws_summary["B3"] = "Value"
    ws_summary["A3"].fill = HEADER_FILL
    ws_summary["B3"].fill = HEADER_FILL
    ws_summary["A3"].font = HEADER_FONT
    ws_summary["B3"].font = HEADER_FONT

    labels = {
        "annualised_return": "Annualised Return",
        "sharpe_ratio": "Sharpe Ratio (annualised, rf=0%)",
        "win_rate": "Win Rate",
        "max_drawdown": "Max Drawdown",
        "num_round_trips": "Number of Round-Trip Trades",
        "total_pnl": "Total P&L (GBP)",
        "final_portfolio_value": "Final Portfolio Value (GBP)",
        "num_pairs_tested": "Pairs Tested",
        "num_pairs_cointegrated": "Cointegrated Pairs (ADF p<0.05)",
        "num_pairs_traded": "Pairs Traded",
    }
    row = 4
    for key, label in labels.items():
        if key not in summary_metrics:
            continue
        value = summary_metrics[key]
        if key in ("annualised_return", "win_rate", "max_drawdown"):
            value = f"{value:.2%}" if pd.notna(value) else "n/a"
        elif key in ("sharpe_ratio",):
            value = f"{value:.2f}" if pd.notna(value) else "n/a"
        elif key in ("total_pnl", "final_portfolio_value"):
            value = f"{value:,.2f}"
        ws_summary.cell(row=row, column=1, value=label)
        ws_summary.cell(row=row, column=2, value=value)
        row += 1

    ws_summary.column_dimensions["A"].width = 34
    ws_summary.column_dimensions["B"].width = 20

    if "portfolio_curve" in chart_paths:
        img = XLImage(chart_paths["portfolio_curve"])
        img.anchor = "D3"
        ws_summary.add_image(img)
    if "drawdown_curve" in chart_paths:
        img2 = XLImage(chart_paths["drawdown_curve"])
        img2.anchor = "D28"
        ws_summary.add_image(img2)

    # --- Pair Stats sheet ----------------------------------------------------
    ws_stats = wb.create_sheet("Pair Stats")
    _write_dataframe(ws_stats, pair_stats)

    # --- Trade Log sheet -----------------------------------------------------
    ws_trades = wb.create_sheet("Trade Log")
    if trade_log is not None and not trade_log.empty:
        _write_dataframe(ws_trades, trade_log)
    else:
        ws_trades["A1"] = "No trades executed."

    # --- Charts sheet (per-pair thumbnails) -----------------------------------
    ws_charts = wb.create_sheet("Pair Charts")
    anchor_row = 1
    for pair_name, path in chart_paths.get("pair_charts", {}).items():
        ws_charts.cell(row=anchor_row, column=1, value=pair_name).font = Font(bold=True)
        img = XLImage(path)
        img.width, img.height = 480, 430
        ws_charts.add_image(img, f"A{anchor_row + 1}")
        anchor_row += 24

    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log.info("Excel dashboard written to %s", out_path)
